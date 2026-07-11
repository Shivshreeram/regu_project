"""Diagnose document extraction and translation issues."""

from docx import Document

print("[ORIGINAL DOCUMENT STRUCTURE]")
print("=" * 80)

# Load original
doc_orig = Document('xarelto_first_3_pages.docx')
print(f"\nOriginal paragraphs: {len(doc_orig.paragraphs)}")
print(f"Original tables: {len(doc_orig.tables)}")

print("\n[FIRST 5 PARAGRAPHS]")
for i, para in enumerate(doc_orig.paragraphs[:5]):
    text = para.text.strip()
    if text:
        preview = text[:80] + "..." if len(text) > 80 else text
        print(f"  Para {i}: {preview}")

if doc_orig.tables:
    print(f"\n[TABLE 0 - First 3 rows]")
    table = doc_orig.tables[0]
    print(f"  Rows: {len(table.rows)}, Cols: {len(table.columns)}")
    for r_idx, row in enumerate(table.rows[:3]):
        cells_preview = []
        for cell in row.cells:
            cell_text = cell.text.strip()[:25]
            cells_preview.append(cell_text)
        print(f"    Row {r_idx}: {cells_preview}")

# Load translated
print("\n" + "=" * 80)
print("[TRANSLATED DOCUMENT STRUCTURE]")
print("=" * 80)

doc_trans = Document('xarelto_first_3_pages_Translated_DE_v2.docx')
print(f"\nTranslated paragraphs: {len(doc_trans.paragraphs)}")
print(f"Translated tables: {len(doc_trans.tables)}")

print("\n[FIRST 5 TRANSLATED PARAGRAPHS]")
for i, para in enumerate(doc_trans.paragraphs[:5]):
    text = para.text.strip()
    if text:
        preview = text[:80] + "..." if len(text) > 80 else text
        print(f"  Para {i}: {preview}")

if doc_trans.tables:
    print(f"\n[TRANSLATED TABLE 0 - First 3 rows]")
    table = doc_trans.tables[0]
    print(f"  Rows: {len(table.rows)}, Cols: {len(table.columns)}")
    for r_idx, row in enumerate(table.rows[:3]):
        cells_preview = []
        for cell in row.cells:
            cell_text = cell.text.strip()[:25]
            cells_preview.append(cell_text)
        print(f"    Row {r_idx}: {cells_preview}")

print("\n" + "=" * 80)
